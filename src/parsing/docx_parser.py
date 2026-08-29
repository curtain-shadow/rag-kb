# -*- coding: utf-8 -*-
"""DOCX 解析（python-docx）：抽取段落 + 表格，表格按行拼接成文本。"""
from pathlib import Path

import docx

from .base import BaseParser, ParsedDocument, make_doc_id, register


@register
class DocxParser(BaseParser):
    exts = (".docx",)

    def parse(self, path: Path) -> ParsedDocument:
        d = docx.Document(str(path))
        sections: list[str] = []

        # 正文段落
        for p in d.paragraphs:
            t = p.text.strip()
            if t:
                sections.append(t)

        # 表格：每行用 " | " 拼接各单元格，保留表格语义
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    sections.append(" | ".join(cells))

        return ParsedDocument(
            doc_id=make_doc_id(str(path)),
            source=str(path),
            title=path.stem,
            sections=sections,
            text="\n\n".join(sections),
            metadata={"paragraphs": len(d.paragraphs), "tables": len(d.tables), "format": "docx"},
        )
